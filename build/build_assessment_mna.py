#!/usr/bin/env python3
"""Build the v2.0 WA and Case Study assessment set for TGS-2026061581.

The item count, instrument type and K/A mappings preserve the approved legacy
assessment. The case and model responses are rebuilt against the v2.0 course.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import course_data as C

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
OUT = ROOT / "assessment"
LOGO = ROOT / "courseware" / "assets" / "tertiary-infotech-logo.png"
BLUE = RGBColor(0x1F, 0x6F, 0xEB)
INK = RGBColor(0x11, 0x18, 0x27)
GREY = RGBColor(0x55, 0x5B, 0x66)
RED = RGBColor(0xB9, 0x1C, 0x1C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def shade(cell, fill):
    prop = cell._tc.get_or_add_tcPr()
    tag = OxmlElement("w:shd")
    tag.set(qn("w:fill"), fill)
    prop.append(tag)


def new_doc():
    d = Document()
    sec = d.sections[0]
    sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75); sec.right_margin = Inches(0.75)
    st = d.styles["Normal"]
    st.font.name = "Arial"; st.font.size = Pt(10.5); st.font.color.rgb = INK
    for name, size, color in [("Title", 24, BLUE), ("Heading 1", 16, BLUE),
                              ("Heading 2", 13, INK), ("Heading 3", 11, BLUE)]:
        s = d.styles[name]; s.font.name = "Arial"; s.font.size = Pt(size)
        s.font.bold = True; s.font.color.rgb = color
    return d


def para(d, text="", bold=False, size=10.5, color=INK, after=5, align=None):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    if align is not None: p.alignment = align
    r = p.add_run(text); r.bold = bold; r.font.name = "Arial"
    r.font.size = Pt(size); r.font.color.rgb = color
    return p


def segments(d, bits, after=5):
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
    for text, bold, color in bits:
        r = p.add_run(text); r.bold = bold; r.font.name = "Arial"
        r.font.size = Pt(10.5); r.font.color.rgb = color
    return p


def bullet(d, text, level=0):
    p = d.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def number(d, text):
    p = d.add_paragraph(style="List Number"); p.paragraph_format.space_after = Pt(4)
    p.add_run(text); return p


def page(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def footer(d, confidential=False):
    p = d.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = "TRAINER CONFIDENTIAL · " if confidential else ""
    r = p.add_run(f"{label}{C.COURSE_CODE} · Version {C.DOC_VERSION} · Tertiary Infotech Academy Pte Ltd")
    r.font.name = "Arial"; r.font.size = Pt(8); r.font.color.rgb = RED if confidential else GREY


def cover(d, kind, answer_key=False):
    d.add_paragraph()
    if LOGO.exists():
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO), width=Inches(2.15))
    para(d, "TERTIARY INFOTECH ACADEMY PTE LTD", True, 13, INK, 2, WD_ALIGN_PARAGRAPH.CENTER)
    para(d, "UEN: 201200696W", False, 9, GREY, 26, WD_ALIGN_PARAGRAPH.CENTER)
    para(d, kind.upper(), True, 23, BLUE if not answer_key else RED, 10, WD_ALIGN_PARAGRAPH.CENTER)
    if answer_key:
        para(d, "ANSWER KEY & MARKING GUIDE", True, 16, RED, 14, WD_ALIGN_PARAGRAPH.CENTER)
        para(d, "TRAINER / ASSESSOR CONFIDENTIAL", True, 11, RED, 7, WD_ALIGN_PARAGRAPH.CENTER)
    para(d, C.TITLE, True, 18, INK, 10, WD_ALIGN_PARAGRAPH.CENTER)
    para(d, f"Course Code: {C.COURSE_CODE}", True, 11, BLUE, 3, WD_ALIGN_PARAGRAPH.CENTER)
    para(d, f"TSC: {C.TSC_CODE} — {C.TSC_TITLE}", False, 10, GREY, 18, WD_ALIGN_PARAGRAPH.CENTER)
    para(d, f"Version {C.DOC_VERSION} · 13 August 2026", True, 11, BLUE, 2, WD_ALIGN_PARAGRAPH.CENTER)
    if answer_key:
        para(d, "Not for candidate distribution", True, 10, RED, 2, WD_ALIGN_PARAGRAPH.CENTER)
    page(d)


def cell_text(cell, text, bold=False, color=INK, size=9.5):
    cell.text = ""; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    r = cell.paragraphs[0].add_run(str(text)); r.bold = bold; r.font.name = "Arial"
    r.font.size = Pt(size); r.font.color.rgb = color


def instructions(d, instrument, duration, item_count, evidence):
    d.add_heading("Trainee Information and Assessment Instructions", 1)
    t = d.add_table(rows=4, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, label in enumerate(["Candidate name (as per NRIC/FIN)", "Last 3 digits + letter of NRIC/FIN", "Assessment date", "Assessor name"]):
        cell_text(t.cell(row, 0), label, True); shade(t.cell(row, 0), "E8F0FE")
        cell_text(t.cell(row, 1), "")
    d.add_heading("Instructions to Candidate", 2)
    for text in [
        f"This is an individual, open-book {instrument} assessment.",
        f"Time allowed: {duration}. Answer all {item_count} questions/tasks.",
        "Use your own words. State reasonable assumptions and show workings where applicable.",
        "Base your response only on the stated case facts, course assumptions and evidence supplied in the question paper.",
        "Do not invent transaction terms or present course assumptions as reported facts.",
        f"Submit the completed paper through the LMS at {C.LMS_URL} and follow the assessor's submission instructions.",
    ]: number(d, text)
    d.add_heading("Competency and Grading", 2)
    para(d, evidence)
    g = d.add_table(rows=3, cols=4); g.style = "Table Grid"
    headers = ["Result", "C / NYC", "Assessor signature", "Date"]
    for i, h in enumerate(headers): cell_text(g.cell(0, i), h, True, WHITE, 9); shade(g.cell(0, i), "1F6FEB")
    for row, label in [(1, "First attempt"), (2, "Reassessment, if applicable")]:
        cell_text(g.cell(row, 0), label, True)
        for col in range(1, 4): cell_text(g.cell(row, col), "")
    para(d, "C = Competent · NYC = Not Yet Competent", False, 8.5, GREY, 0)
    page(d)


def answer_space(d, lines=5):
    para(d, "Candidate response:", True, 9.5, GREY, 2)
    for _ in range(lines): para(d, "________________________________________________________________________________", False, 9, GREY, 2)


WA = [
    {
        "map": "K1", "slides": "Topic 1; Activity 1",
        "q": "Bonadio announced that T. Gschwender & Associates (TGA) would join its Internal Audit unit, with transaction terms undisclosed. Explain three transaction instruments or structures a buyer could evaluate for this combination. For each, state one reason it may fit and one risk or trade-off.",
        "answer": [
            "A full share or asset acquisition can deliver control and operating integration, but requires careful liability, consent and tax analysis.",
            "An earn-out or deferred consideration can bridge uncertainty about client retention and future performance, but creates measurement and behaviour disputes unless definitions and governance are precise.",
            "Seller rollover equity can align key owners with post-close value creation and retention, but adds governance, liquidity and exit-right complexity.",
            "Other defensible instruments taught in Topic 1 earn credit when suitability and trade-off are correctly linked to the case.",
        ],
    },
    {
        "map": "K2, K5", "slides": "Topic 2; Activities 3–4",
        "q": "Describe how a buyer should forecast the financial and tax implications of the proposed combination over three years. Your answer must distinguish reported facts from assumptions and explain at least four checks that protect the decision from overstated value.",
        "answer": [
            "Build a standalone baseline for both firms, then model revenue and cost synergies separately with timing, probability, owner and cost-to-achieve.",
            "Translate assumptions into P&L, cash flow and balance-sheet effects, including purchase consideration, transaction/integration costs, working capital, financing and accounting impacts.",
            "Assess tax attributes and leakage: deal form, deductibility, indirect/transfer taxes, jurisdiction, limits on interest or loss use, and uncertainty requiring specialist advice.",
            "Key controls include scenario/sensitivity analysis, no double counting, finance-independent challenge, reconciliation to the deal thesis, and clear fact/assumption/source labels.",
        ],
    },
    {
        "map": "K4", "slides": "Topic 3; Activities 5–6",
        "q": "Explain four target-industry risks that are especially relevant when a professional-services firm acquires a credit-risk and loan-review specialist. For each risk, name evidence you would seek during due diligence.",
        "answer": [
            "Client concentration and relationship portability: client-level revenue, churn, engagement terms, consent/change-of-control clauses and pipeline quality.",
            "Key-person and talent risk: partner/employee revenue attribution, tenure, restrictive covenants, compensation, succession and retention sentiment.",
            "Regulatory/professional liability: licences, regulator correspondence, methodology, independence/conflicts, complaints, claims and insurance.",
            "Data, cyber and technology risk: data inventory, access controls, incidents, vendor dependencies, model governance and privacy obligations.",
            "Other supported industry risks can earn credit, including reputation, integration capacity and economic/credit-cycle exposure.",
        ],
    },
    {
        "map": "K2", "slides": "Topic 4; Activities 7–8",
        "q": "Explain how culture, talent, client continuity and integration governance can affect the value of a professional-services acquisition. Recommend four practical measures that protect the deal thesis from signing through post-close integration.",
        "answer": [
            "Value resides in people, relationships, methods and reputation; unwanted departures or client disruption can remove the revenue assumed in valuation.",
            "Use explicit integration principles, a SteerCo/IMO/workstream structure, decision rights and an issue-escalation path.",
            "Build a talent/retention plan using role clarity, listening, career paths and appropriately designed rollover/earn-out or retention mechanisms.",
            "Segment clients, assign owners, prepare continuity messages and track retention/relationship signals.",
            "Maintain a value-capture register linking each synergy to an owner, milestone, baseline, cash impact and cost-to-achieve.",
        ],
    },
]


CASE_FACTS = [
    "On 6 August 2026, Bonadio announced that TGA was expected to join its Internal Audit unit in its Advisory & Consulting practice; the report stated an expected 16 August 2026 close.",
    "TGA, founded in 1984, provides credit-risk management, loan review and regulatory-advisory services to financial institutions.",
    "Transaction terms were not disclosed in the cited report. The assessment therefore uses the explicit course assumptions below; they are not reported deal facts.",
]

CASE_ASSUMPTIONS = [
    "TGA training-case revenue: US$12.0m; EBITDA: US$2.4m; debt-free/cash-free basis; normalised working capital is met at close.",
    "Illustrative valuation evidence: EV/EBITDA comparable range 7.0x–9.0x; precedent range 7.5x–9.5x; DCF enterprise value range US$18.5m–US$23.5m.",
    "Base-case synergies: US$0.6m, US$1.4m and US$2.0m in Years 1–3; cost-to-achieve: US$1.2m in Year 1 and US$0.4m in Year 2. Apply a 25% downside to synergies for a downside case.",
    "Illustrative risk evidence: top three clients contribute 38% of revenue; 45% of revenue is linked to five senior professionals; two clients require change-of-control consent; one open cyber-remediation item is rated high; no quantified loss has been provided.",
]

CS = [
    {
        "map": "A1", "slides": "Activities 1–2",
        "task": "Prepare an indicative transaction recommendation. Triangulate the three valuation methods into a defensible enterprise-value range, recommend a transaction structure, and state the assumptions and protections required before proceeding.",
        "answer": [
            "Calculate the comparable range as US$16.8m–US$21.6m (7.0x–9.0x × US$2.4m) and the precedent range as US$18.0m–US$22.8m (7.5x–9.5x × US$2.4m); retain the stated DCF range of US$18.5m–US$23.5m.",
            "Recommend a reasoned overlap/weighting range rather than averaging mechanically; for example US$18.5m–US$22.0m, subject to diligence and buyer-specific synergy policy.",
            "A defensible structure may combine cash at close with deferred/earn-out consideration and/or seller rollover linked to client and key-talent retention.",
            "State that price/terms are illustrative and undisclosed; identify conditions such as consent completion, working-capital mechanism, warranties/indemnities, retention arrangements and cyber remediation.",
        ],
    },
    {
        "map": "A1, A2", "slides": "Activities 3–4",
        "task": "Build the three-year synergy and pro-forma challenge. Show base and downside synergy, cost-to-achieve, net synergy by year, and at least five financial/tax validation checks. Explain which findings could change price, terms or the decision.",
        "answer": [
            "Base net synergy: Year 1 = -US$0.6m; Year 2 = US$1.0m; Year 3 = US$2.0m. Cumulative three-year net = US$2.4m.",
            "Downside gross synergy at 75%: US$0.45m, US$1.05m, US$1.50m. After stated cost-to-achieve: -US$0.75m, US$0.65m, US$1.50m; cumulative = US$1.40m.",
            "Checks should include baseline integrity, timing/ramp, probability, owner/action, double counting, cost-to-achieve, cash conversion, working capital, financing, purchase accounting and reconciliation to the investment case.",
            "Tax analysis should challenge deal form, deductibility, indirect/transfer taxes, jurisdiction, loss/interest limitations and adviser confirmation.",
            "If downside value cannot support price, or unresolved tax/cash leakage is material, reduce price, defer/condition consideration, add protection, remediate pre-close or stop.",
        ],
    },
    {
        "map": "A2, A3, A4, A5, A6", "slides": "Activities 5–8",
        "task": "Submit a one-page deal-team recommendation covering: target/diligence conclusions; a prioritised risk register and business-intelligence evidence plan; stakeholder alignment; a Day One/integration governance design; and a go/conditional-go/no-go recommendation with explicit decision gates.",
        "answer": [
            "Prioritise client-consent/concentration, key-person dependency and high cyber remediation; show likelihood, impact, evidence gap, mitigation, owner, timing and residual risk.",
            "Use business intelligence to triangulate internal financial/client/people data with contracts, regulator/professional evidence, market/competitor information and independent confirmations; document source, date, reliability and limitations.",
            "Stakeholder plan should cover owners, leaders, critical talent, clients, regulators, vendors and integration teams, with owner, message, channel, timing and feedback/escalation route.",
            "Design SteerCo, IMO and functional workstreams; define pre-close legal interaction rules, Day One must-haves, operating-model decisions, integration milestones and a value-capture register.",
            "A strong answer is likely conditional-go: secure required client consents and key-person arrangements, validate cyber scope/cost, confirm valuation/downside support and approve integration capacity before final gates. A different recommendation earns credit if evidence-led and internally consistent.",
        ],
    },
]


def build_wa_paper():
    d = new_doc(); footer(d)
    cover(d, "Written Assessment (WA-SAQ)")
    instructions(d, "written", "1 hour", "4", "All four responses must demonstrate satisfactory knowledge across K1, K2, K4 and K5. The assessor records one overall C/NYC outcome.")
    d.add_heading("Written Assessment Questions", 1)
    para(d, "Case anchor: Use the announced Bonadio–TGA transaction only where the question refers to it. The report says terms were undisclosed.", False, 9.5, GREY, 10)
    for i, item in enumerate(WA, 1):
        d.add_heading(f"Question {i}: {item['map']}", 2)
        para(d, item["q"], False, 10.5, INK, 5)
        answer_space(d, 14)
        if i != len(WA): page(d)
    path = OUT / f"WA (SAQ) - Strategic Mergers and Acquisitions - {C.VERSION}.docx"
    d.save(path); return path


def build_wa_key():
    d = new_doc(); footer(d, True)
    cover(d, "Written Assessment (WA-SAQ)", True)
    d.add_heading("Assessor Guidance", 1)
    para(d, "Use professional judgement. Accept equivalent evidence-led answers aligned to the course. Do not require identical wording. Record clear remediation where evidence is insufficient.")
    for i, item in enumerate(WA, 1):
        if i > 1: page(d)
        d.add_heading(f"Question {i}: {item['map']}", 1)
        para(d, item["q"], False, 10, GREY, 8)
        para(d, "Suggestive answers / model evidence", True, 11, BLUE, 4)
        for point in item["answer"]: bullet(d, point)
        para(d, f"Course traceability: {item['slides']} · Activity folder(s) as cited.", False, 9, GREY, 6)
        para(d, "Assessor decision:  ☐ Satisfactory   ☐ Not yet satisfactory", True, 10, RED, 2)
    path = OUT / f"Answer to WA (SAQ) - Strategic Mergers and Acquisitions - {C.VERSION}.docx"
    d.save(path); return path


def case_block(d):
    d.add_heading("Case Study Scenario: Bonadio–TGA Announced Transaction", 1)
    para(d, "Reported facts", True, 11, BLUE, 3)
    for x in CASE_FACTS: bullet(d, x)
    para(d, "Course-only assumptions for assessment", True, 11, RED, 3)
    for x in CASE_ASSUMPTIONS: bullet(d, x)
    para(d, "Use the reported facts and assumptions exactly as labelled. You may make further reasonable assumptions if you state and justify them.", True, 9.5, GREY, 5)


def build_cs_paper():
    d = new_doc(); footer(d)
    cover(d, "Case Study Assessment (CS)")
    instructions(d, "case-study", "1 hour", "3", "All three tasks must be satisfactory. Together they assess A1–A6. The assessor records one overall C/NYC outcome.")
    case_block(d)
    for i, item in enumerate(CS, 1):
        page(d)
        d.add_heading(f"Task {i}: {item['map']}", 1)
        para(d, item["task"], False, 10.5, INK, 8)
        answer_space(d, 11)
    path = OUT / f"CS Assessment - Strategic Mergers and Acquisitions - {C.VERSION}.docx"
    d.save(path); return path


def build_cs_key():
    d = new_doc(); footer(d, True)
    cover(d, "Case Study Assessment (CS)", True)
    d.add_heading("Assessor Guidance", 1)
    para(d, "The model evidence demonstrates one competent route. Accept alternative calculations, structures and recommendations where workings are correct, assumptions are explicit, and the conclusion follows the evidence.")
    case_block(d)
    for i, item in enumerate(CS, 1):
        page(d)
        d.add_heading(f"Task {i}: {item['map']}", 1)
        para(d, item["task"], False, 10, GREY, 8)
        para(d, "Suggestive answers / model evidence", True, 11, BLUE, 4)
        for point in item["answer"]: bullet(d, point)
        para(d, f"Course traceability: {item['slides']} and the corresponding detailed activity folders.", False, 9, GREY, 6)
        para(d, "Assessor decision:  ☐ Satisfactory   ☐ Not yet satisfactory", True, 10, RED, 2)
    path = OUT / f"Answer to CS Assessment - Strategic Mergers and Acquisitions - {C.VERSION}.docx"
    d.save(path); return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for built in [build_wa_paper(), build_wa_key(), build_cs_paper(), build_cs_key()]:
        print(f"Saved {built}")
