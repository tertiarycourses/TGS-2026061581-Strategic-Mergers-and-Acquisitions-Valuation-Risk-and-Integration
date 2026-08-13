#!/usr/bin/env python3
"""Build aligned activity folders, Learner Guide and Lesson Plan."""
import json, os, re, sys, textwrap
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent; sys.path.insert(0,str(HERE))
import course_data as C
from data_domain1 import DOMAIN1
from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3
from data_domain4 import DOMAIN4
import prodoc

ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4
REPO=HERE.parent
CW=REPO/"courseware"; ACTIVITIES=REPO/"activities"; ASSETS=CW/"assets"
BRAND=RGBColor(0x1F,0x6F,0xEB); TEAL=RGBColor(0x10,0xB9,0x81)
DARK=RGBColor(0x11,0x18,0x27); GREY=RGBColor(0x55,0x5B,0x66)

def slug(s):
    return re.sub(r"[^a-z0-9]+","-",s.lower()).strip("-")

def write_activity_files():
    ACTIVITIES.mkdir(parents=True,exist_ok=True)
    index=[f"# Activities — {C.TITLE}","",f"Course: {C.COURSE_CODE}","",
           "The eight activities form one progressive real-world case around the announced Bonadio–TGA transaction. Reported facts are separated from course-only assumptions.",""]
    for a in ACT:
        folder=ACTIVITIES/f"activity-{a['num']:02d}-{slug(a['title'])}"
        folder.mkdir(parents=True,exist_ok=True)
        lines=[f"# Activity {a['num']}: {a['title']}","",f"**TSC mapping:** {a['criteria']}  ",f"**Objective:** {a['objective']}  ",f"**Expected output:** {a['output']}  ",f"**Primary source:** {a['source']}","",
               "## Goal","",a['desc'],"","## Real-world scenario","",a['fact_basis'],"",a['scenario'],"",
               "## Evidence boundary","","Reported deal facts and learner assumptions must never be mixed. Financial terms not disclosed by the parties are treated only as training assumptions.","",
               "### Training assumptions / guardrails",""]
        lines += [f"- {x}" for x in a["assumptions"]]
        lines += ["","## Questions",""]+[f"{i}. {q}" for i,q in enumerate(a["questions"],1)]
        lines += ["","## Detailed activity procedure",""]
        for i,(name,detail) in enumerate(a["steps"],1):
            lines += [f"### Step {i}: {name}","",detail,"",f"**Checkpoint:** Record the evidence, calculation or decision produced in this step.",""]
        lines += ["## Deliverable","",a["output"],"","## Acceptance criteria","",a["test"],"",
                  "## Reflection and debrief questions","",
                  "1. Which assumption changed your recommendation most, and why?",
                  "2. What evidence would you request before relying on the result?",
                  "3. Which stakeholder owns the next decision or mitigation?",
                  "4. How does this activity feed the next stage of the M&A lifecycle?","",
                  "## Source note","",f"Source accessed 13 August 2026: {a['source']}",""]
        (folder/"README.md").write_text("\n".join(lines),encoding="utf-8")
        template=[f"# Submission Template — Activity {a['num']}","",f"## {a['title']}","",
                  "### 1. Reported facts","","- ","","### 2. Training assumptions","","- ","",
                  "### 3. Analysis and calculations","","","### 4. Answers to activity questions",""]
        template += [f"#### Question {i}\n\n{q}\n\n**Response:**\n" for i,q in enumerate(a["questions"],1)]
        template += ["","### 5. Recommendation","","","### 6. Risks, mitigations and owner","","",
                     "### 7. Evidence and sources","","- ","","### 8. Acceptance check","",f"- [ ] {a['test']}",""]
        (folder/"submission-template.md").write_text("\n".join(template),encoding="utf-8")
        index.append(f"- [Activity {a['num']}: {a['title']}](activity-{a['num']:02d}-{slug(a['title'])}/README.md) — {a['criteria']}")
    (ACTIVITIES/"README.md").write_text("\n".join(index)+"\n",encoding="utf-8")

def base_doc(kind):
    d=Document(); sec=d.sections[0]
    sec.top_margin=Inches(0.75); sec.bottom_margin=Inches(0.75); sec.left_margin=Inches(0.8); sec.right_margin=Inches(0.8)
    st=d.styles["Normal"]; st.font.name="Arial"; st.font.size=Pt(10.5)
    st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.08
    prodoc.style_headings(d)
    prodoc.add_cover_page(d,kind,C.TITLE,C.DOC_VERSION,org_logo=str(ASSETS/"tertiary-infotech-logo.png"),course_logo=None,course_code=C.COURSE_CODE)
    prodoc.add_version_control(d,[(C.DOC_VERSION,C.VERSION_DATE,
        "Expanded source-aligned release with real-world Bonadio–TGA activities, valuation, risk and integration content.",C.TRAINER)])
    prodoc.add_toc(d); return d

def add_hyperlink(paragraph, url, text=None):
    part=paragraph.part; rid=part.relate_to(url,"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
    link=OxmlElement("w:hyperlink"); link.set(qn("r:id"),rid)
    run=OxmlElement("w:r"); rpr=OxmlElement("w:rPr"); color=OxmlElement("w:color"); color.set(qn("w:val"),"1F6FEB")
    underline=OxmlElement("w:u"); underline.set(qn("w:val"),"single"); rpr.extend([color,underline]); run.append(rpr)
    t=OxmlElement("w:t"); t.text=text or url; run.append(t); link.append(run); paragraph._p.append(link)

def bullet(doc,text,level=0):
    p=doc.add_paragraph(style="List Bullet" if level==0 else "List Bullet 2"); p.add_run(text); return p

def number(doc,text):
    # Use explicit numbering in the supplied text. Word's shared List Number
    # counter otherwise continues across unrelated sections and can turn an
    # activity's Question 1 into item 11 or 120 after rendering.
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(0.22)
    p.paragraph_format.first_line_indent=Inches(-0.22); p.add_run(text); return p

def note(doc,label,text,color=BRAND):
    table=doc.add_table(rows=1,cols=1); table.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=table.cell(0,0); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; prodoc._shade_cell(c,"F5F8FC")
    p=c.paragraphs[0]; r=p.add_run(label+": "); r.bold=True; r.font.color.rgb=color; p.add_run(text)
    doc.add_paragraph()

def set_cell(cell,text,bold=False,fill=None,color=DARK,size=9.2):
    cell.text=""; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill: prodoc._shade_cell(cell,fill)
    p=cell.paragraphs[0]; r=p.add_run(str(text)); r.bold=bold; r.font.name="Arial"; r.font.size=Pt(size); r.font.color.rgb=color

def topic_slide_ranges(slide_map,total=999):
    starts={int(k):v for k,v in slide_map["topics"].items()}; out={}
    nums=sorted(starts)
    for i,n in enumerate(nums): out[n]=(starts[n], starts[nums[i+1]]-1 if i+1<len(nums) else total)
    return out

def build_lesson_plan():
    d=base_doc("LESSON PLAN")
    slide_map=json.loads((CW/"slide_map.json").read_text())
    from pptx import Presentation
    slide_total=len(Presentation(str(next(CW.glob(f"{C.SHORT_TITLE}-*.pptx")))).slides)
    ranges=topic_slide_ranges(slide_map,slide_total)
    d.add_heading("Course Overview",1)
    d.add_paragraph("This two-day WSQ course develops the ability to structure, value, risk-manage and integrate M&A transactions. A progressive real-world case links strategy, finance, diligence, governance and post-merger execution.")
    t=d.add_table(rows=0,cols=2); t.style="Table Grid"
    for k,v in [("Course Code",C.COURSE_CODE),("TSC",f"{C.TSC_CODE} — {C.TSC_TITLE}"),("Duration","2 days · 16 hours"),("Training","14 hours"),("Assessment","WA 1 hour + Case Study 1 hour"),("Daily Time","9:00am–6:00pm")]:
        c=t.add_row().cells; set_cell(c[0],k,True,"E8F0FE"); set_cell(c[1],v)
    d.add_heading("Learning Outcomes",1)
    for x in C.LEARNING_OUTCOMES: bullet(d,x)
    d.add_heading("Daily Schedule",1)
    for day,rows in C.SCHEDULE.items():
        d.add_heading(f"Day {day} — {C.DAY_THEMES[day]}",2)
        table=d.add_table(rows=1,cols=5); table.style="Table Grid"; table.alignment=WD_TABLE_ALIGNMENT.CENTER
        for i,h in enumerate(["Time","Minutes","Topic / Activity","Method","Slides"]): set_cell(table.rows[0].cells[i],h,True,"1F6FEB",RGBColor(255,255,255),9)
        total=0
        for start,end,mins,kind,text in rows:
            c=table.add_row().cells; fill="FDE9D9" if kind=="lunch" else ("FFF4E5" if kind=="break" else ("E8F7EE" if kind=="assess" else None))
            slide_ref="—"
            topic_matches=[n for n in ranges if f"Topic {n}" in text]
            if topic_matches:
                n=topic_matches[0]; slide_ref=f"Slides {ranges[n][0]}–{ranges[n][1]}"
            elif "Activities" in text:
                ns=[int(x) for x in re.findall(r"\d+",text.split("Activities",1)[1])]
                if ns: slide_ref=", ".join(f"A{n}: {slide_map['activities'].get(str(n),'?')}" for n in range(min(ns),max(ns)+1))
            method={"admin":"Interactive presentation","topic":"Interactive presentation · case study · discussion","break":"Break","lunch":"Lunch","recap":"Reflection · Q&A","assess":"Assessment"}[kind]
            for i,v in enumerate([f"{start}–{end}",mins,text,method,slide_ref]): set_cell(c[i],v,kind=="assess",fill,size=8.6)
            if kind!="lunch": total+=mins
        p=d.add_paragraph(); r=p.add_run(f"Total training/assessment time excluding lunch: {total} minutes ({total/60:.1f} hours)."); r.italic=True; r.font.color.rgb=GREY
        assert total==480,(day,total)
    d.add_heading("Topic and Activity Breakdown",1)
    for topic in C.TOPICS:
        lo,hi=ranges[topic["num"]]
        d.add_heading(f"Topic {topic['num']} — {topic['title']} · Slides {lo}–{hi}",2)
        d.add_paragraph(f"Mapping: {topic['criteria']}. {topic['subtitle']}.")
        for a in [x for x in ACT if x["topic"]==topic["num"]]:
            d.add_heading(f"Activity {a['num']} — {a['title']} · Slide {slide_map['activities'][str(a['num'])]}",3)
            d.add_paragraph(a["desc"]); bullet(d,f"Output: {a['output']}"); bullet(d,f"Acceptance: {a['test']}")
    d.add_heading("Resources Required",1)
    for x in ["Trainer slide deck and Learner Guide","Individual activity folders and submission templates","Spreadsheet software for Activities 2–4","Internet access for cited public sources","LMS access for courseware and assessment"]: bullet(d,x)
    d.add_heading("Assessment",1)
    for x in [C.ASSESSMENT["written"],C.ASSESSMENT["practical"],C.ASSESSMENT["note"],"Assessment block: Day 2, 4:00pm–6:00pm.","Order: SSG attendance → WA → CS → LMS submission → Assessment Summary Record."]: bullet(d,x)
    p=d.add_paragraph("LMS: "); add_hyperlink(p,C.LMS_URL)
    prodoc.add_page_numbers(d); prodoc.enable_update_fields(d)
    out=CW/f"LP-{C.SHORT_TITLE}.docx"; d.save(out); print("Saved",out)

def build_learner_guide():
    d=base_doc("LEARNER GUIDE")
    d.add_heading("How to Use This Guide",1)
    d.add_paragraph("Read the concept section before each activity, then complete the activity in its individual folder. Detailed procedures belong here and in the activity Markdown files; the slide deck remains visual and concept-led.")
    note(d,"Evidence boundary","The Bonadio–TGA transaction is real and sourced. Because financial terms were not disclosed, all numeric case inputs are explicitly labelled training assumptions and must never be described as transaction facts.")
    d.add_heading("Course Learning Outcomes and TSC Mapping",1)
    for x in C.LEARNING_OUTCOMES: bullet(d,x)
    d.add_heading("Knowledge",2)
    for x in C.KNOWLEDGE: bullet(d,x)
    d.add_heading("Abilities",2)
    for x in C.ABILITIES: bullet(d,x)
    d.add_heading("The Progressive Real-World Case",1)
    d.add_paragraph(DOMAIN1[0]["fact_basis"])
    p=d.add_paragraph("Primary transaction report: "); add_hyperlink(p,DOMAIN1[0]["source"])
    d.add_paragraph("Activities 1–8 follow the same deal from strategic rationale and valuation through forecast, diligence, risk, integration, culture and client continuity. This continuity allows each output to become an input to the next decision.")
    for topic in C.TOPICS:
        d.add_page_break(); d.add_heading(f"Topic {topic['num']} — {topic['title']}",1)
        d.add_paragraph(f"TSC mapping: {topic['criteria']}. {topic['subtitle']}.")
        d.add_heading("Key Concepts",2)
        for h,body in topic["concepts"]: bullet(d,f"{h}: {body}")
        for item in topic["lessons"]:
            d.add_heading(item[1],2)
            if item[0] in ("tile","cards"):
                for h,body in item[2]: bullet(d,f"{h}: {body}")
            elif item[0]=="flow":
                for i,x in enumerate(item[2],1): number(d,f"{i}. {x}")
            elif item[0]=="two":
                d.add_heading("Perspective A",3)
                for x in item[2]: bullet(d,x)
                d.add_heading("Perspective B",3)
                for x in item[3]: bullet(d,x)
            elif item[0]=="big": note(d,"Key point",item[1]+" "+item[2],TEAL)
        for a in [x for x in ACT if x["topic"]==topic["num"]]:
            d.add_page_break(); d.add_heading(f"Activity {a['num']} — {a['title']}",1)
            d.add_paragraph(f"TSC mapping: {a['criteria']}. {a['objective']}.")
            d.add_heading("Goal and Deliverable",2); d.add_paragraph(a["desc"]); bullet(d,f"Deliverable: {a['output']}")
            d.add_heading("Real-World Scenario",2); d.add_paragraph(a["fact_basis"]); d.add_paragraph(a["scenario"])
            p=d.add_paragraph("Source: "); add_hyperlink(p,a["source"])
            d.add_heading("Training Assumptions and Guardrails",2)
            for x in a["assumptions"]: bullet(d,x)
            d.add_heading("Scenario Questions",2)
            for i,q in enumerate(a["questions"],1): number(d,f"Question {i}: {q}")
            d.add_heading("Detailed Step-by-Step Activity",2)
            for i,(name,detail) in enumerate(a["steps"],1):
                d.add_heading(f"Step {i} — {name}",3); d.add_paragraph(detail); note(d,"Checkpoint","Record the evidence, calculation or decision produced before moving to the next step.")
            d.add_heading("Acceptance Test",2); note(d,"Pass condition",a["test"],TEAL)
            d.add_heading("Reflection and Debrief",2)
            for i,q in enumerate(["Which assumption changed your recommendation most, and why?","What additional evidence would you request?","Which stakeholder owns the next action?","How does this output feed the next M&A stage?"],1): number(d,f"{i}. {q}")
    d.add_heading("Assessment Preparation",1)
    for x in [C.ASSESSMENT["written"],C.ASSESSMENT["practical"],"Review the activity acceptance tests and be able to explain your judgement in your own words.","Only learner-facing question papers are opened on the LMS; answer keys are trainer-only."]: bullet(d,x)
    d.add_heading("Assessment Flow",2)
    for i,x in enumerate(["Complete SSG/TRAQOM assessment digital attendance.","Sit the WA-SAQ, then the Case Study.","Upload completed answers to the LMS.","Sign the Assessment Summary Record."],1): number(d,f"{i}. {x}")
    d.add_heading("Sources and Further Reading",1)
    for n,u in C.SOURCES:
        p=d.add_paragraph(); p.add_run(n+": ").bold=True; add_hyperlink(p,u)
    d.add_heading("Glossary",1)
    glossary=[("Enterprise value","Value of operations available to debt and equity capital providers."),("Equity value","Residual value attributable to equity after the enterprise-value bridge."),("Normalised EBITDA","Earnings adjusted to a maintainable, comparable operating level."),("WACC","Weighted average cost of capital used to discount enterprise cash flows."),("Synergy","Incremental value available because the businesses combine."),("Cost-to-achieve","One-off spending required to realise a synergy."),("Quality of earnings","Diligence analysis of sustainability, cash conversion and accounting quality."),("IMO","Integration Management Office coordinating workstreams, risks, decisions and value."),("Day One","First day of legal ownership/control after closing."),("TSA","Transitional Services Agreement supporting temporary post-close services.")]
    for term,definition in glossary: bullet(d,f"{term}: {definition}")
    prodoc.add_page_numbers(d); prodoc.enable_update_fields(d)
    out=CW/f"LG-{C.SHORT_TITLE}.docx"; d.save(out); print("Saved",out)

if __name__=="__main__":
    write_activity_files(); build_lesson_plan(); build_learner_guide()
