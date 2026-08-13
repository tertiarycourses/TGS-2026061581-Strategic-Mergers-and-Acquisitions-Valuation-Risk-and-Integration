#!/usr/bin/env python3
"""Generate the TGS-2026061581 course slide deck (all-white Tertiary house style).

Design helpers are the same set used by the tertiary-course-slides skill that
produced the n8n reference deck (cover, section, content, two_col, cards3,
big_statement, step_slide, test_slide, brk). Content is driven entirely by
course_data.py + data_domainN.py so the deck stays 100% aligned with the LP,
LG and activities.
"""
import os, sys, json
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import course_data as C
from data_domain1 import DOMAIN1
from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3
from data_domain4 import DOMAIN4
ACTIVITIES = DOMAIN1 + DOMAIN2 + DOMAIN3 + DOMAIN4

REPO = os.path.dirname(HERE)
ASSETS = os.path.join(REPO, "courseware", "assets")
ORIGINAL_VISUALS = os.path.join(HERE, "original_visuals")

ORIGINAL_GALLERIES = {
    1: [
        ("How consideration moves", [("s025-cash-consideration.png", "Cash consideration"), ("s027-escrow-holdback.png", "Escrow / holdback"), ("s031-convertible-instrument.png", "Convertible instrument")]),
        ("How risk is allocated", [("s037-contingent-instruments.png", "Contingent consideration"), ("s051-asset-purchase.png", "Asset purchase"), ("s053-retention-incentives.png", "Retention mechanisms")]),
        ("How value is triangulated", [("s057-valuation-methods.png", "Valuation families"), ("s061-income-valuation.png", "Income approach"), ("s071-hybrid-valuation.png", "Hybrid triangulation")]),
    ],
    2: [
        ("How value grows", [("s075-value-levers.png", "Value levers"), ("s087-shareholder-value-roadmap.png", "Value roadmap"), ("s090-economies-of-scale.png", "Scale pathways")]),
        ("How forecasts connect", [("s103-forecasting-impact.png", "Business-impact forecast"), ("s104-forecasting-prism.png", "Forecast dimensions"), ("s120-mna-analysis.png", "Integrated analysis")]),
        ("How finance is challenged", [("s121-tax-considerations.png", "Tax considerations"), ("s143-financial-ratios.png", "Financial ratios"), ("s144-cash-conversion-cycle.jpg", "Cash conversion cycle")]),
    ],
    3: [
        ("How targets are screened", [("s153-target-screening.png", "Target screening"), ("s162-due-diligence.png", "Due diligence"), ("s164-diligence-workstreams.png", "Diligence workstreams")]),
        ("How risk becomes evidence", [("s168-risk-map.png", "Risk map"), ("s169-business-intelligence.png", "Business intelligence"), ("s170-intelligence-flow.png", "Intelligence flow")]),
        ("How the decision is framed", [("s171-pestle.png", "PESTLE lens"), ("s176-evaluation-tools.png", "Evaluation toolkit"), ("s177-evaluation-tree.png", "Evaluation tree")]),
    ],
    4: [
        ("How stakeholders govern", [("s183-stakeholder-engagement.jpg", "Stakeholder engagement"), ("s189-checks-balances.png", "Checks and balances"), ("s200-culture-change.png", "Culture and change")]),
        ("How integration is mobilised", [("s184-post-merger-integration.jpg", "Post-merger integration"), ("s193-implementation-plan.png", "Implementation plan"), ("s200-culture-change.png", "Culture and change")]),
        ("How the original story connects", [("s074-valuation-strategy.png", "Valuation strategy"), ("s088-scale-strategy.png", "Scale strategy"), ("s145-transaction-examples.png", "Transaction examples")]),
    ],
}

# ---------------- palette (matches reference) ----------------
BLUE=RGBColor(0x1F,0x6F,0xEB); TEAL=RGBColor(0x10,0xB9,0x81); AMBER=RGBColor(0xF5,0x9E,0x0B)
INK=RGBColor(0x16,0x1B,0x26); GREY=RGBColor(0x5B,0x63,0x72); LIGHT=RGBColor(0xF5,0xF8,0xFC)
WHITE=RGBColor(0xFF,0xFF,0xFF); LINE=RGBColor(0xE2,0xE8,0xF0); VIOLET=RGBColor(0x7C,0x3A,0xED)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height
BLANK=prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)
def rect(s,x,y,w,h,color,line=None):
    sp=s.shapes.add_shape(1,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(1)
    sp.shadow.inherit=False; return sp
def oval(s,x,y,w,h,color):
    sp=s.shapes.add_shape(9,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    sp.line.fill.background(); sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,space=4):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,line in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(space)
        for t,sz,col,bold in line:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=bold
            r.font.color.rgb=col; r.font.name="Arial"
    return tb
def bullets(s,x,y,w,h,items,size=18,color=INK,gap=10,mcolor=BLUE):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(gap)
        lvl=it[1] if isinstance(it,tuple) else 0
        text=it[0] if isinstance(it,tuple) else it
        r=p.add_run(); r.text=("•  " if lvl==0 else "–  ")+text
        r.font.size=Pt(size if lvl==0 else size-2); r.font.color.rgb=color if lvl==0 else GREY
        r.font.name="Arial"; r.font.bold=(lvl==0 and isinstance(it,tuple) and len(it)>2 and it[2])
    return tb

PAGE={"n":0}
def footer(s):
    PAGE["n"]+=1
    txt(s,Inches(0.4),Inches(7.05),Inches(7.5),Inches(0.35),
        [[(f"{C.SHORT_TITLE}  ·  {C.COURSE_CODE}",9,GREY,False)]])
    txt(s,Inches(5.0),Inches(7.05),Inches(3.3),Inches(0.35),
        [[("© 2026 Tertiary Infotech Academy Pte Ltd",9,GREY,False)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(12.4),Inches(7.05),Inches(0.6),Inches(0.35),
        [[(str(PAGE["n"]),9,GREY,False)]],align=PP_ALIGN.RIGHT)
def head(s,title,kicker=None,kcolor=BLUE):
    rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),Inches(1.55),kcolor)
    if kicker: txt(s,Inches(0.85),Inches(0.5),Inches(11.6),Inches(0.4),[[(kicker,14,kcolor,True)]])
    txt(s,Inches(0.85),Inches(0.9),Inches(11.9),Inches(0.9),[[(title,29,INK,True)]])
    rect(s,Inches(0.85),Inches(1.7),Inches(11.63),Inches(0.02),LINE)
    return s
def _logo(name):
    p=os.path.join(ASSETS,name)
    return p if os.path.exists(p) else None

# ---------------- slide templates ----------------
def cover():
    s=slide(); rect(s,0,0,SW,SH,WHITE)
    rect(s,0,0,SW,Inches(0.22),BLUE); rect(s,0,Inches(7.28),SW,Inches(0.22),TEAL)
    org=_logo("tertiary-infotech-logo.png")
    if org: s.shapes.add_picture(org,Inches(0.85),Inches(0.7),height=Inches(1.05))
    # course badge (top-right)
    rect(s,Inches(10.75),Inches(0.72),Inches(1.8),Inches(1.0),BLUE)
    txt(s,Inches(10.75),Inches(0.82),Inches(1.8),Inches(0.5),[[("M&A",24,WHITE,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(10.75),Inches(1.28),Inches(1.8),Inches(0.4),[[("VALUE · RISK · PMI",8,WHITE,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(0.9),Inches(2.3),Inches(12),Inches(0.6),[[("COURSE SLIDES  ·  WSQ",16,BLUE,True)]])
    txt(s,Inches(0.9),Inches(2.85),Inches(12.0),Inches(1.9),[[(C.TITLE,40,INK,True)]])
    rect(s,Inches(0.92),Inches(4.75),Inches(2.4),Inches(0.06),TEAL)
    txt(s,Inches(0.9),Inches(5.05),Inches(12),Inches(1.4),
        [[(f"WSQ Course Code: {C.COURSE_CODE}",16,GREY,False)],
         [("Conducted by Tertiary Infotech Academy Pte Ltd  ·  UEN 201200696W",14,GREY,False)]],space=6)
    txt(s,Inches(0.9),Inches(6.5),Inches(12),Inches(0.4),[[(f"Version {C.VERSION}  ·  {C.VERSION_DATE}",12,GREY,False)]])
    txt(s,Inches(0.9),Inches(6.85),Inches(12),Inches(0.34),[[("© 2026 Tertiary Infotech Academy Pte Ltd. All rights reserved.  ·  www.tertiarycourses.com.sg",10,GREY,False)]])

def section(kicker,title,n,sub=""):
    s=slide(); rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),SH,BLUE)
    rect(s,Inches(0.85),Inches(2.5),Inches(0.14),Inches(2.0),TEAL)
    txt(s,Inches(1.25),Inches(2.55),Inches(11),Inches(0.6),[[(kicker,18,BLUE,True)]])
    txt(s,Inches(1.25),Inches(3.0),Inches(11.4),Inches(1.6),[[(title,40,INK,True)]])
    if sub: txt(s,Inches(1.27),Inches(4.55),Inches(11),Inches(0.8),[[(sub,16,GREY,False)]])
    txt(s,Inches(10.0),Inches(0.7),Inches(2.8),Inches(1.6),[[(n,72,RGBColor(0xE2,0xE8,0xF0),True)]],align=PP_ALIGN.RIGHT)
    footer(s)
def content(title,items,kicker=None,size=20):
    s=head(slide(),title,kicker); bullets(s,Inches(0.85),Inches(1.95),Inches(11.6),Inches(4.9),items,size=size); footer(s); return s
def two_col(title,left,right,kicker=None,lhead="",rhead=""):
    s=head(slide(),title,kicker)
    rect(s,Inches(0.85),Inches(1.95),Inches(5.7),Inches(4.7),LIGHT); rect(s,Inches(6.95),Inches(1.95),Inches(5.55),Inches(4.7),LIGHT)
    if lhead: txt(s,Inches(1.1),Inches(2.15),Inches(5.2),Inches(0.4),[[(lhead,16,BLUE,True)]])
    if rhead: txt(s,Inches(7.2),Inches(2.15),Inches(5.0),Inches(0.4),[[(rhead,16,TEAL,True)]])
    bullets(s,Inches(1.1),Inches(2.7),Inches(5.2),Inches(3.8),left,size=16)
    bullets(s,Inches(7.2),Inches(2.7),Inches(5.05),Inches(3.8),right,size=16,mcolor=TEAL); footer(s); return s
def cards3(title,cards,kicker):
    s=head(slide(),title,kicker); xs=[Inches(0.85),Inches(5.0),Inches(9.15)]
    for i,c in enumerate(cards[:3]):
        x=xs[i]; col=c[0]
        rect(s,x,Inches(1.95),Inches(3.65),Inches(4.7),LIGHT); rect(s,x,Inches(1.95),Inches(3.65),Inches(0.12),col)
        txt(s,x+Inches(0.25),Inches(2.2),Inches(3.2),Inches(0.6),[[(c[1],19,col,True)]])
        bullets(s,x+Inches(0.25),Inches(2.95),Inches(3.2),Inches(3.4),c[2],size=14,mcolor=col,gap=9)
    footer(s); return s
def big_statement(line1,line2,kicker,color=BLUE):
    s=slide(); rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),SH,color)
    txt(s,Inches(1.1),Inches(2.2),Inches(11),Inches(0.5),[[(kicker,16,color,True)]])
    txt(s,Inches(1.1),Inches(2.8),Inches(11.3),Inches(2.4),[[(line1,38,INK,True)]])
    if line2: txt(s,Inches(1.12),Inches(4.9),Inches(11),Inches(1.2),[[(line2,20,GREY,False)]])
    footer(s); return s
import math
PALETTE=[BLUE,TEAL,VIOLET,AMBER]
def tile_grid(title,items,kicker=None,cols=2,size=15,icons=None,accent=BLUE):
    """Grid of light panels, each with a coloured icon/number badge + text.
    items: list of strings (or (title,caption) tuples). Much richer than a bullet list."""
    s=head(slide(),title,kicker,kcolor=accent)
    n=len(items); rows=math.ceil(n/cols)
    X0=Inches(0.85); Y0=Inches(1.95); TOTW=Inches(11.63); AREAH=Inches(4.78)
    gx=Inches(0.3); gy=Inches(0.26)
    cw=int((TOTW-gx*(cols-1))/cols); ch=int((AREAH-gy*(rows-1))/rows)
    bd=Inches(0.6)
    for i,it in enumerate(items):
        r=i//cols; c=i%cols
        x=int(X0+(cw+gx)*c); y=int(Y0+(ch+gy)*r); col=PALETTE[i%len(PALETTE)]
        rect(s,x,y,cw,ch,LIGHT); rect(s,x,y,Inches(0.1),ch,col)
        oval(s,x+Inches(0.28),int(y+ch/2-bd/2),bd,bd,col)
        ic=icons[i] if icons else str(i+1)
        txt(s,x+Inches(0.28),int(y+ch/2-bd/2),bd,bd,[[(ic,19,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        tx=x+Inches(1.08); tw=cw-Inches(1.32)
        if isinstance(it,tuple):
            txt(s,tx,int(y+Inches(0.14)),tw,int(ch-Inches(0.2)),
                [[(it[0],size+2,INK,True)],[(it[1],size-2,GREY,False)]],anchor=MSO_ANCHOR.MIDDLE,space=3)
        else:
            txt(s,tx,int(y+Inches(0.1)),tw,int(ch-Inches(0.16)),[[(it,size,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s
def original_gallery(title, entries, topic_no):
    """Three-image gallery using diagrams extracted from the supplied old PPT."""
    s=head(slide(),title,"ORIGINAL DECK · VISUAL REFERENCE",kcolor=VIOLET)
    xs=[Inches(0.75),Inches(4.58),Inches(8.41)]
    for i,(filename,caption) in enumerate(entries):
        x=xs[i]; y=Inches(1.85); w=Inches(3.55); h=Inches(4.65)
        rect(s,x,y,w,h,WHITE,line=LINE); rect(s,x,y,w,Inches(0.1),PALETTE[i])
        path=os.path.join(ORIGINAL_VISUALS,filename)
        with Image.open(path) as im: iw,ih=im.size
        maxw,maxh=Inches(3.15),Inches(3.45); scale=min(maxw/iw,maxh/ih)
        pw,ph=int(iw*scale),int(ih*scale)
        px=int(x+(w-pw)/2); py=int(y+Inches(0.32)+(maxh-ph)/2)
        s.shapes.add_picture(path,px,py,width=pw,height=ph)
        txt(s,x+Inches(0.14),Inches(5.65),w-Inches(0.28),Inches(0.55),[[(caption,15,INK,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    txt(s,Inches(0.8),Inches(6.58),Inches(11.8),Inches(0.28),[[("Reused from the supplied original PPT; interpret using the updated v2.0 concepts and evidence.",9,GREY,False)]],align=PP_ALIGN.CENTER)
    footer(s); return s
def flow_h(title,steps,kicker=None,color=BLUE):
    """Horizontal numbered flow: coloured chips connected by chevrons."""
    s=head(slide(),title,kicker,kcolor=color)
    n=len(steps); X0=Inches(0.85); TOTW=Inches(11.63); gap=Inches(0.34)
    cw=int((TOTW-gap*(n-1))/n); y=Inches(2.55); ch=Inches(3.15); bd=Inches(0.82)
    for i,st in enumerate(steps):
        x=int(X0+(cw+gap)*i)
        rect(s,x,y,cw,ch,LIGHT); rect(s,x,y,cw,Inches(0.1),color)
        oval(s,int(x+cw/2-bd/2),int(y+Inches(0.42)),bd,bd,color)
        txt(s,int(x+cw/2-bd/2),int(y+Inches(0.42)),bd,bd,[[(str(i+1),30,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        txt(s,x+Inches(0.16),int(y+Inches(1.55)),cw-Inches(0.32),int(ch-Inches(1.7)),[[(st,14,INK,False)]],align=PP_ALIGN.CENTER)
        if i<n-1:
            txt(s,int(x+cw-Inches(0.04)),int(y+ch/2-Inches(0.3)),int(gap+Inches(0.08)),Inches(0.6),
                [[("▶",15,color,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s
def trainer_slide(kicker,name,role,rows,initials,accent=BLUE):
    """Profile-card layout: avatar badge + name/role panel on the left, labelled
    info tiles on the right. rows: list of (LABEL, value); blank value → fill-in line."""
    s=head(slide(),"About the Trainer",kicker,kcolor=accent)
    lx=Inches(0.85); lw=Inches(3.65)
    rect(s,lx,Inches(1.95),lw,Inches(4.7),LIGHT); rect(s,lx,Inches(1.95),lw,Inches(0.12),accent)
    bd=Inches(1.7); ax=int(lx+(lw-bd)/2)
    oval(s,ax,Inches(2.5),bd,bd,accent)
    txt(s,ax,Inches(2.5),bd,bd,[[(initials,44,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    txt(s,lx+Inches(0.15),Inches(4.55),lw-Inches(0.3),Inches(0.6),[[(name,21,INK,True)]],align=PP_ALIGN.CENTER)
    txt(s,lx+Inches(0.15),Inches(5.2),lw-Inches(0.3),Inches(1.2),[[(role,13,GREY,False)]],align=PP_ALIGN.CENTER)
    rx=Inches(4.9); rw=Inches(7.6); ry=Inches(1.95); rh=Inches(4.7)
    n=len(rows); gy=Inches(0.2); th=int((rh-gy*(n-1))/n)
    for i,(label,val) in enumerate(rows):
        y=int(ry+(th+gy)*i); col=PALETTE[i%len(PALETTE)]
        rect(s,rx,y,rw,th,LIGHT); rect(s,rx,y,Inches(0.1),th,col)
        vruns=[(val,14,INK,False)] if val else [("____________________________________________",13,LINE,False)]
        txt(s,rx+Inches(0.32),y,rw-Inches(0.6),th,
            [[(label.upper(),11,col,True)],vruns],anchor=MSO_ANCHOR.MIDDLE,space=3)
    footer(s); return s
def activity_overview(tag,title,desc,build,services,kicker):
    s=head(slide(),title,kicker,kcolor=TEAL)
    rect(s,Inches(0.85),Inches(1.85),Inches(1.7),Inches(0.5),TEAL)
    txt(s,Inches(0.85),Inches(1.9),Inches(1.7),Inches(0.4),[[(tag,16,WHITE,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(0.85),Inches(2.55),Inches(11.7),Inches(1.6),[[(desc,21,INK,False)]])
    rect(s,Inches(0.85),Inches(4.3),Inches(11.7),Inches(2.0),LIGHT)
    txt(s,Inches(1.1),Inches(4.5),Inches(11),Inches(0.4),[[("You'll build",14,BLUE,True)]])
    txt(s,Inches(1.1),Inches(4.9),Inches(11),Inches(0.6),[[(build,18,INK,True)]])
    txt(s,Inches(1.1),Inches(5.6),Inches(11.2),Inches(0.6),[[("Tools / evidence:  ",13,GREY,True),(services,13,GREY,False)]]); footer(s); return s
def step_slide(kicker,act_title,n,total,text,cmd=""):
    s=head(slide(),act_title,kicker,TEAL)
    oval(s,Inches(0.85),Inches(2.5),Inches(1.4),Inches(1.4),TEAL)
    txt(s,Inches(0.85),Inches(2.74),Inches(1.4),Inches(0.9),[[(str(n),38,WHITE,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(0.95),Inches(1.95),Inches(11),Inches(0.4),[[(f"STEP {n} OF {total}",13,GREY,True)]])
    txt(s,Inches(2.55),Inches(2.4),Inches(10.1),Inches(1.3),[[(text,23,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
    if cmd:
        rect(s,Inches(2.55),Inches(4.15),Inches(10.1),Inches(0.95),RGBColor(0x0B,0x12,0x20))
        txt(s,Inches(2.8),Inches(4.28),Inches(9.7),Inches(0.7),[[("$ "+cmd,13,RGBColor(0x9C,0xDC,0xFE),False)]],anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s
def test_slide(act_title,text,kicker):
    s=head(slide(),act_title,kicker,TEAL)
    rect(s,Inches(0.85),Inches(2.3),Inches(11.7),Inches(2.6),RGBColor(0xE8,0xF7,0xEE))
    txt(s,Inches(1.2),Inches(2.6),Inches(11),Inches(0.5),[[("✅  Test it",20,RGBColor(0x12,0x7A,0x3E),True)]])
    txt(s,Inches(1.2),Inches(3.3),Inches(11),Inches(1.4),[[(text,18,INK,False)]]); footer(s); return s
def brk(kind,dur,color=AMBER):
    s=slide(); rect(s,0,0,SW,SH,WHITE)
    rect(s,0,0,SW,Inches(0.22),color); rect(s,0,Inches(7.28),SW,Inches(0.22),color)
    rect(s,Inches(5.4),Inches(2.35),Inches(2.53),Inches(0.1),color)
    txt(s,0,Inches(2.75),SW,Inches(1.2),[[(kind,48,INK,True)]],align=PP_ALIGN.CENTER)
    txt(s,0,Inches(4.05),SW,Inches(0.8),[[(dur,22,color,True)]],align=PP_ALIGN.CENTER); PAGE["n"]+=1

# ============================================================ COURSE BUILD
def link_slide(title, url, cards, kicker):
    """Visual access slide with a real clickable URL."""
    s=head(slide(),title,kicker,kcolor=VIOLET)
    x=Inches(0.85); y=Inches(2.0); w=Inches(11.65); h=Inches(0.72)
    rect(s,x,y,w,h,RGBColor(0xEE,0xE9,0xFE))
    box=txt(s,x+Inches(0.25),y+Inches(0.12),w-Inches(0.5),Inches(0.45),[[(url,15,VIOLET,True)]],align=PP_ALIGN.CENTER)
    for p in box.text_frame.paragraphs:
        for r in p.runs: r.hyperlink.address=url
    xs=[Inches(0.85),Inches(4.85),Inches(8.85)]
    for i,(label,body) in enumerate(cards[:3]):
        rect(s,xs[i],Inches(3.15),Inches(3.65),Inches(2.75),LIGHT)
        rect(s,xs[i],Inches(3.15),Inches(3.65),Inches(0.11),PALETTE[i])
        txt(s,xs[i]+Inches(0.22),Inches(3.45),Inches(3.2),Inches(0.55),[[(label,18,PALETTE[i],True)]])
        txt(s,xs[i]+Inches(0.22),Inches(4.15),Inches(3.2),Inches(1.45),[[(body,15,INK,False)]])
    footer(s); return s

def render_lesson(topic, item):
    kind=item[0]; title=item[1]; kicker=f"TOPIC {topic['code']} · {topic['criteria']}"
    if kind=="tile": return tile_grid(title,item[2],kicker=kicker,cols=2,size=14)
    if kind=="flow":
        # Dense nine-step frameworks need a 3×3 map; narrow horizontal cards
        # force mid-word line breaks and undermine training-room readability.
        if len(item[2]) > 6:
            return tile_grid(title,[(f"Step {i+1}",text) for i,text in enumerate(item[2])],kicker=kicker,cols=3,size=12,accent=BLUE)
        return flow_h(title,item[2],kicker=kicker,color=BLUE)
    if kind=="cards":
        cols=[BLUE,TEAL,VIOLET]
        cards=[(cols[i%3],x[0],[x[1]]) for i,x in enumerate(item[2])]
        return cards3(title,cards,kicker=kicker)
    if kind=="two": return two_col(title,item[2],item[3],kicker=kicker,lhead="Perspective A",rhead="Perspective B")
    if kind=="big": return big_statement(item[1],item[2],kicker,color=BLUE)
    raise ValueError(f"Unknown lesson type: {kind}")

def add_transitions(presentation, kind="fade", speed="med"):
    for sl in presentation.slides:
        el=sl._element
        for tr in el.findall(qn("p:transition")): el.remove(tr)
        tr=el.makeelement(qn("p:transition"), {"spd":speed})
        tr.append(el.makeelement(qn(f"p:{kind}"), {})); el.append(tr)

SLIDE_MAP={"topics":{},"activities":{}}
cover()
section("COURSE ADMINISTRATION","Welcome, Outcomes and Assessment","00")
tile_grid("Digital Attendance (Mandatory)",[
 ("AM, PM and Assessment","Complete every SSG digital-attendance event shown by the trainer."),
 ("Scan the QR code","Use the mobile camera and submit the generated attendance form."),
 ("Verify submission","Tell the trainer immediately if the confirmation does not appear."),
 ("Funding condition","At least 75% attendance and a Competent result are required.")],kicker="TRAQOM · SSG DIGITAL ATTENDANCE",cols=2,size=14)
trainer_slide("YOUR TRAINER · GENERAL","Your Trainer","General Trainer template — to be completed by the trainer",
 [("Name",""),("Title / Designation",""),("Qualifications",""),("Areas of expertise",""),("Training & industry experience",""),("Contact","")],initials="?",accent=GREY)
trainer_slide("YOUR TRAINER",C.TRAINER,"Principal Trainer · Tertiary Infotech Academy Pte Ltd",
 [("Role","Principal Trainer and course facilitator"),("Specialisms","Business, technology and applied decision-making"),("Delivery","WSQ professional development and workplace application"),("Contact","enquiry@tertiaryinfotech.com")],initials="AA",accent=BLUE)
tile_grid("Let's Know Each Other",[
 ("Role","Your organisation, function and decision responsibilities."),
 ("M&A exposure","A deal, diligence, valuation or integration experience."),
 ("Learning goal","The decision or capability you want to strengthen."),
 ("Case lens","One stakeholder whose perspective you will represent today.")],kicker="INTRODUCTION",cols=2,size=14)
tile_grid("Ground Rules",[
 ("Participate","Challenge assumptions and contribute evidence."),("Respect","Disagree with ideas, not people."),
 ("Confidentiality","Use public or course data only; do not disclose live deal information."),("Devices","Silent mode; step out quietly if needed."),
 ("Time","Return from breaks promptly."),("Attendance","Maintain the required SSG attendance record.")],kicker="HOUSEKEEPING",cols=2,size=14)
two_col("Course Outline — 2 Days",[
 ("Day 1 — Structure, Value and Forecast",0),
 ("Topic 1: deal instruments, valuation and scale",1),
 ("Activities 1–2: thesis/structure and valuation",1),
 ("Topic 2: forecast, synergies, statements and tax",1),
 ("Activities 3–4: synergy model and board challenge",1)],
 [("Day 2 — Risk, Intelligence and Integration",0),
 ("Topic 3: screening, diligence, risk and BI",1),
 ("Activities 5–6: diligence and risk intelligence",1),
 ("Topic 4: governance, people and integration",1),
 ("Activities 7–8; WA 4–5pm; CS 5–6pm",1)],kicker="LESSON PLAN",lhead="Day 1",rhead="Day 2")
tile_grid("Learning Outcomes",[(f"LO{i+1}",x.split(": ",1)[1]) for i,x in enumerate(C.LEARNING_OUTCOMES)],kicker="WHAT YOU WILL ACHIEVE",cols=2,size=13)
cards3("Skills Framework Alignment",[
 (BLUE,C.TSC_CODE,[C.TSC_TITLE,"One 5-level Technical Skill and Competency"]),
 (TEAL,"Knowledge K1–K5",["Instruments · forecasts · valuation · industry · financial/tax impact"]),
 (VIOLET,"Abilities A1–A6",["Value · controls · integration · risk · intelligence · opportunity evaluation"])],kicker="TSC ALIGNMENT")
tile_grid("Briefing for Assessment",[
 ("Individual","No discussion, photography or recording."),("Open book","Use slides, Learner Guide and approved course materials only."),
 ("Submission","Answer in your own words and upload the completed document to the LMS."),("Conduct","Phones and other materials under the table; follow assessor instructions."),
 ("Outcome","Competent or Not Yet Competent; reassessment and appeal are available."),("Sign-off","Complete digital attendance and sign the Assessment Summary Record.")],kicker="BEFORE ASSESSMENT",cols=2,size=13)
cards3("Assessment and Funding",[
 (BLUE,"WA-SAQ · 1 hour",["4 open-ended questions","Covers K1–K5","Open book"]),
 (TEAL,"Case Study · 1 hour",["3 integrated tasks","Covers A1–A6","Open book"]),
 (VIOLET,"Competency",["Satisfactory evidence","At least 75% attendance","Appeal process available"])],kicker="FINAL ASSESSMENT")
flow_h("Assessment Flow",["TRAQOM / SSG attendance","Assessment digital attendance","Sit WA then Case Study","Submit answers on the LMS","Sign the Assessment Summary Record"],kicker="ON ASSESSMENT DAY",color=VIOLET)
link_slide("Access the Course Activities",C.GITHUB_URL,[
 ("Clone","git clone the repository, then open the activities folder."),
 ("Download ZIP","Use GitHub Code → Download ZIP when Git is unavailable."),
 ("One folder each","Every activity includes scenario, questions, steps and acceptance criteria.")],kicker="HANDS-ON MATERIALS")
link_slide("Download Course Material",C.LMS_URL,[
 ("1 · Sign in","Open the LMS with your registered email and password or OTP."),
 ("2 · Open course","Select this TGS course from your enrolled courses."),
 ("3 · Download","Download slides and the Learner Guide; assessments are completed on the LMS.")],kicker="COURSEWARE AND ASSESSMENT")

section("CORE MODEL","One Deal · One Evidence Chain","C")
flow_h("The decision chain",["Strategic thesis", "Valuation and structure", "Diligence and risk", "Approval and close", "Integration and value"],kicker="END-TO-END VIEW",color=TEAL)
tile_grid("Evidence hierarchy",[
 ("Verified public fact","Published or regulatory information with source and date."),
 ("Management assertion","A claim to be tested against documents and operating evidence."),
 ("Training assumption","A simulated input used only for learning and clearly labelled."),
 ("Independent evidence","Third-party, customer, market, legal or specialist verification.")],kicker="DO NOT CONFUSE FACT WITH ASSUMPTION",cols=2,size=14)
big_statement("A strong M&A recommendation links thesis → evidence → value → risk → owner.","If one link is missing, the decision is not ready.","COURSE CASE",color=TEAL)

TOPIC_ACTS={t["num"]:[a for a in ACTIVITIES if a["topic"]==t["num"]] for t in C.TOPICS}
for t in C.TOPICS:
    SLIDE_MAP["topics"][str(t["num"])]=len(prs.slides)+1
    section(f"TOPIC {t['code']} · {t['criteria']}",t["title"],t["code"],t["subtitle"])
    tile_grid(f"Key Concepts — {t['title']}",t["concepts"],kicker=f"TSC {C.TSC_CODE}",cols=2,size=14)
    for item in t["lessons"]: render_lesson(t,item)
    for gallery_title, gallery_entries in ORIGINAL_GALLERIES[t["num"]]:
        original_gallery(gallery_title,gallery_entries,t["num"])
    acts=TOPIC_ACTS[t["num"]]
    cards3(f"Applied Activities — {t['title']}",[
        (BLUE,f"Activity {acts[0]['num']}",[acts[0]["title"],acts[0]["criteria"]]),
        (TEAL,f"Activity {acts[1]['num']}",[acts[1]["title"],acts[1]["criteria"]]),
        (VIOLET,"Progressive case",["Bonadio–TGA announced transaction","Facts separated from course assumptions"])],kicker="REAL-WORLD CASE")
    for a in acts:
        SLIDE_MAP["activities"][str(a["num"])]=len(prs.slides)+1
        activity_overview(f"ACTIVITY {a['num']}",a["title"],a["desc"],a["build"],a["services"],kicker=f"TOPIC {t['code']} · {a['criteria']}")
        two_col("Case Scenario — Reported Facts and Guardrails",
            [a["fact_basis"],f"Source: {a['source']}"],a["assumptions"][:3],
            kicker=f"ACTIVITY {a['num']} · REAL TRANSACTION",lhead="Reported / sourced",rhead="Training assumptions / limits")
        tile_grid("Activity Questions",a["questions"],kicker=f"ACTIVITY {a['num']} · DISCUSS AND DECIDE",cols=2,size=13,accent=VIOLET)
        test_slide(a["title"],a["test"],kicker=f"ACTIVITY {a['num']} · ACCEPTANCE TEST")
    tile_grid(f"Recap — {t['title']}",[(a["title"],a["objective"]) for a in acts]+[
        ("Decision habit","Separate facts, assumptions and unresolved evidence."),
        ("Transfer","Carry valuation and diligence commitments into integration ownership.")],kicker="TOPIC RECAP",cols=2,size=13)
    if t["num"]==1: brk("Lunch Break","60 minutes",color=AMBER)
    if t["num"]==2: brk("End of Day 1","Return tomorrow ready to challenge the deal",color=BLUE)
    if t["num"]==3: brk("Lunch Break","60 minutes",color=AMBER)

section("WRAP-UP","From Deal Thesis to Realised Value","W")
tile_grid("What You Achieved",[(f"LO{i+1}",x.split(": ",1)[1]) for i,x in enumerate(C.LEARNING_OUTCOMES)],kicker="COURSE SUMMARY",cols=2,size=13)
flow_h("The complete M&A discipline",["Thesis", "Value and structure", "Diligence", "Risk-adjusted decision", "Integration and benefits"],kicker="TAKEAWAY",color=TEAL)
tile_grid("Evidence Sources",[(n,u) for n,u in C.SOURCES[:4]],kicker="COURSE ENRICHMENT",cols=2,size=11)
tile_grid("Evidence Sources — Continued",[(n,u) for n,u in C.SOURCES[4:]],kicker="COURSE ENRICHMENT",cols=2,size=11)
link_slide("Courseware and Assessment on the LMS",C.LMS_URL,[
 ("Courseware","Download the current slides and Learner Guide."),
 ("Assessment","Open and complete the WA then Case Study."),
 ("Submission","Upload completed answers and complete the sign-off process.")],kicker="LMS")
cards3("Assessment Reminder",[
 (BLUE,"WA-SAQ · 4:00–5:00pm",["4 open-ended questions","K1–K5","Open book"]),
 (TEAL,"Case Study · 5:00–6:00pm",["3 integrated tasks","A1–A6","Open book"]),
 (VIOLET,"Finish",["Submit on LMS","Digital attendance","Sign summary record"])],kicker="FINAL ASSESSMENT")
flow_h("Assessment Flow",["TRAQOM / SSG attendance","Assessment digital attendance","Sit WA then Case Study","Submit answers on the LMS","Sign the Assessment Summary Record"],kicker="ON ASSESSMENT DAY",color=VIOLET)
tile_grid("Digital Attendance (Mandatory)",[
 ("Scan","Use the QR code shown by the trainer."),("Submit","Complete the assessment attendance event."),
 ("Confirm","Verify the submission confirmation."),("Resolve","Tell the trainer immediately if it fails.")],kicker="TRAQOM · SSG DIGITAL ATTENDANCE",cols=2,size=14)
big_statement("Thank You","Make the deal thesis testable, the risk visible and the value accountable.","STRATEGIC M&A",color=TEAL)

add_transitions(prs)
os.makedirs(os.path.join(REPO,"courseware"),exist_ok=True)
OUT=os.path.join(REPO,"courseware",f"{C.SHORT_TITLE}-{C.VERSION}.pptx")
prs.save(OUT)
with open(os.path.join(REPO,"courseware","slide_map.json"),"w",encoding="utf-8") as f:
    json.dump(SLIDE_MAP,f,indent=2)
print(f"Saved {OUT} ({len(prs.slides)} slides)")
sys.exit(0)
