#!/usr/bin/env python3
"""Extract selected visual assets from the supplied legacy deck for reuse.

Only embedded pictures are copied; wording and teaching logic continue to come
from the v2.0 single-source course model.
"""
from pathlib import Path
from pptx import Presentation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SOURCE = next((ROOT / "reference").glob("WSQ - Master Trainer Slides*.pptx"))
OUT = HERE / "original_visuals"

# High-value, non-duplicative visuals spanning all four approved topics.
SELECTED = {
    25: "cash-consideration", 27: "escrow-holdback", 31: "convertible-instrument",
    37: "contingent-instruments", 51: "asset-purchase", 53: "retention-incentives",
    57: "valuation-methods", 61: "income-valuation", 71: "hybrid-valuation",
    75: "value-levers", 87: "shareholder-value-roadmap", 90: "economies-of-scale",
    103: "forecasting-impact", 104: "forecasting-prism", 120: "mna-analysis",
    121: "tax-considerations", 143: "financial-ratios", 144: "cash-conversion-cycle",
    153: "target-screening", 162: "due-diligence", 164: "diligence-workstreams",
    168: "risk-map", 169: "business-intelligence", 170: "intelligence-flow",
    171: "pestle", 176: "evaluation-tools", 177: "evaluation-tree",
    183: "stakeholder-engagement", 184: "post-merger-integration", 189: "checks-balances",
    193: "implementation-plan", 200: "culture-change", 204: "governance-culture-change",
    74: "valuation-strategy", 88: "scale-strategy", 145: "transaction-examples",
}


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation(SOURCE)
    written = []
    for slide_no, stem in SELECTED.items():
        slide = prs.slides[slide_no - 1]
        pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
        if not pictures:
            raise RuntimeError(f"Legacy slide {slide_no} has no embedded picture")
        # Use the largest picture on the slide, excluding tiny decorative marks.
        pic = max(pictures, key=lambda shape: shape.width * shape.height)
        ext = pic.image.ext
        path = OUT / f"s{slide_no:03d}-{stem}.{ext}"
        path.write_bytes(pic.image.blob)
        written.append(path)
    print(f"Extracted {len(written)} original visuals to {OUT}")


if __name__ == "__main__":
    main()
